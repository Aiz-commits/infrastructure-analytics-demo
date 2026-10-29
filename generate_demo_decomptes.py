
from pathlib import Path
from decimal import Decimal, ROUND_HALF_UP
from openpyxl import Workbook

folder = Path(__file__).resolve().parent / "demo_decomptes"
folder.mkdir(exist_ok=True)

monthly_works = [480000, 640000, 810000, 820000]


def money(value):
    return Decimal(str(value)).quantize(
        Decimal("0.01"), rounding=ROUND_HALF_UP
    )


cumulative = Decimal("0.00")

for number, monthly in enumerate(monthly_works, start=1):
    works = money(monthly)
    previous = cumulative
    cumulative = money(cumulative + works)

    retention = money(works * Decimal("0.10"))
    net_ht = money(works - retention)
    vat = money(net_ht * Decimal("0.20"))
    ttc = money(net_ht + vat)

    wb = Workbook()
    ws = wb.active
    ws.title = "DP"

    ws["A1"] = f"DEMO PROJECT — DECOMPTE D{number}"
    ws["A2"] = "SYNTHETIC DATA — FOR DEMONSTRATION ONLY"

    # Layout expected by the financial extractor
    ws["A39"] = "Retenue de Garantie"
    ws["B37"] = float(cumulative)
    ws["C37"] = float(previous) if number > 1 else None
    ws["D37"] = float(works)
    ws["D39"] = float(retention)
    ws["D44"] = float(net_ht)
    ws["D45"] = float(vat)
    ws["D46"] = float(ttc)

    filename = folder / f"DEMO_DP_N{number}.xlsx"
    wb.save(filename)

    print(f"Created: {filename.name}")

print("All four synthetic decomptes created successfully!")
